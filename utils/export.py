from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

def export_csv(
    filepath: Path,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
) -> Path:
    filepath = Path(filepath)
    filepath.parent.mkdir(parents=True, exist_ok=True)
    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(headers)
        for row in rows:
            writer.writerow(row)
    return filepath

def export_excel(
    filepath: Path,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    sheet_name: str = "Rapport",
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    filepath = Path(filepath)
    filepath.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    header_fill = PatternFill("solid", fgColor="0D7377")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r_idx, row in enumerate(rows, 2):
        for c_idx, value in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=value)

    for col in ws.columns:
        max_len = 0
        letter = col[0].column_letter
        for cell in col:
            max_len = max(max_len, len(str(cell.value or "")))
        ws.column_dimensions[letter].width = min(max_len + 2, 40)

    wb.save(filepath)
    return filepath

def export_pdf(
    filepath: Path,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    subtitle: str = "",
) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate,
        Table,
        TableStyle,
        Paragraph,
        Spacer,
    )

    filepath = Path(filepath)
    filepath.parent.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(
        str(filepath),
        pagesize=landscape(A4),
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleTN",
        parent=styles["Heading1"],
        textColor=colors.HexColor("#0D7377"),
        spaceAfter=6,
    )
    meta_style = ParagraphStyle(
        "MetaTN", parent=styles["Normal"], fontSize=9, textColor=colors.grey
    )

    elements = [
        Paragraph("StockTN Desktop", title_style),
        Paragraph(title, styles["Heading2"]),
    ]
    if subtitle:
        elements.append(Paragraph(subtitle, meta_style))
    elements.append(
        Paragraph(
            f"Généré le {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            meta_style,
        )
    )
    elements.append(Spacer(1, 12))

    data = [list(headers)] + [list(r) for r in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D7377")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F7F7")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    return filepath
